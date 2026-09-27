"""Slide 4 'Market & Govt. fit' infographic -> native, editable shapes (1 pt = 1 CSS px, slide 1440x810 pt).
Replaces the dashed placeholder at (986, 478, 406x100) in slide4_feasibility_viability.pptx.

Sources (checked 2026-09-27):
- Mordor Intelligence, India Geospatial Analytics Market: USD 1.58 B (2025), 1.81 B (2026), 3.55 B (2031), 14.43% CAGR 2026-31.
- MHA / PIB: 15th Finance Commission, National Disaster Mitigation Fund Rs 13,693 cr for 2021-22 to 2025-26.
- NDMA / HLC 29 Nov 2024: National Landslide Risk Mitigation Project Rs 1,000 cr, 15 states (NDMF).
- NDMA: GLOF risk mitigation Rs 150 cr, Arunachal, HP, Sikkim, Uttarakhand (NDMF).
- NRSC: NDEM geo-portal, developed by ISRO under NDMA guidance, maintained by NRSC (DMSP).

usage: uv run --with python-pptx python market_infographic.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

D = Path(__file__).parent.parent
OX, OY = 986, 478          # placeholder origin on slide 4
NAVY, GREY, GREEN, BLUE = "0B1F44", "5B6675", "0F7B3D", "0C62C6"
ORANGE, RED, PURPLE, LINE = "E0661A", "C00000", "8064A2", "C2CDDD"
FONT = {500: "Montserrat Medium", 600: "Montserrat SemiBold", 700: "Montserrat", 800: "Montserrat ExtraBold"}


def P(v):
    return Emu(int(round(v * 12700)))


def text(g, x, y, w, h, runs, size, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None, spc=None, wrap=True):
    """runs: list of (text, weight, color[, italic])"""
    tb = g.shapes.add_textbox(P(OX + x), P(OY + y), P(w), P(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    for r in runs:
        t, wgt, col = r[:3]
        for k, line in enumerate(t.split("\n")):
            if k:
                p = tf.add_paragraph()
            p.alignment = align
            p.line_spacing = 0.95
            _run(p, line, wgt, col, r, size, spc)
    tb.name = name or "text – " + "".join(r[0] for r in runs).replace("\n", " ")[:40]
    return tb


def _run(p, t, wgt, col, r, size, spc):
        run = p.add_run()
        run.text = t
        f = run.font
        f.size = P(size)
        f.name = FONT[wgt]
        f.bold = wgt == 700
        f.italic = len(r) > 3 and r[3]
        f.color.rgb = RGBColor.from_string(col)
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.makeelement(qn(tag), {"typeface": FONT[wgt]})
            rPr.append(el)
        if spc:
            rPr.set("spc", str(spc))


def rect(g, x, y, w, h, fill=None, line=None, lw=1.0, radius=None, name="shape", dash=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = g.shapes.add_shape(kind, P(OX + x), P(OY + y), P(w), P(h))
    if radius:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = RGBColor.from_string(line); s.line.width = P(lw)
        if dash:
            s.line.dash_style = dash
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.name = name
    return s


def hline(g, x1, y1, x2, y2, col, lw=0.8, dash=None, name="line"):
    c = g.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, P(OX + x1), P(OY + y1), P(OX + x2), P(OY + y2))
    c.line.color.rgb = RGBColor.from_string(col); c.line.width = P(lw)
    if dash:
        c.line.dash_style = dash
    c.name = name
    return c


def build(shapes):
    g = shapes.add_group_shape()
    g.name = "Market & Govt. fit infographic"

    # header (same style as the column's other headings)
    text(g, 0, 2, 220, 15, [("MARKET & GOVT. FIT", 800, GREEN)], 12, spc=48, wrap=False, name="heading – MARKET & GOVT. FIT")
    text(g, 196, 5, 210, 10, [("Mordor Intelligence 2026 · MHA/PIB · NDMA · NRSC", 500, "9AA4B2", True)], 5.8,
         align=PP_ALIGN.RIGHT, name="sources")

    # ---- left: market KPI + bars
    text(g, 0, 22, 90, 20, [("$1.81 B", 800, NAVY)], 17, wrap=False, name="kpi – $1.81 B")
    text(g, 0, 43, 84, 20, [("India geospatial analytics market, 2026", 500, GREY)], 6.5, name="kpi label")
    rect(g, 0, 69, 80, 14, fill="FFF1E6", line=ORANGE, lw=0.9, radius=7, name="chip – CAGR")
    text(g, 0, 69, 80, 14, [("14.43% CAGR", 700, ORANGE)], 7, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, name="chip text – 14.43% CAGR")

    base, top_h = 88, 50                     # baseline y, height of the tallest bar
    bars = [("2025", 1.58, LINE, NAVY), ("2026", 1.81, BLUE, BLUE), ("2031F", 3.55, GREEN, GREEN)]
    for i, (yr, v, fill, lab) in enumerate(bars):
        x, w = 94 + i * 30, 22
        h = top_h * v / 3.55
        rect(g, x, base - h, w, h, fill=fill, name=f"bar – {yr} ${v} B")
        text(g, x - 6, base - h - 9, w + 12, 8, [(f"${v:.2f}B", 700, lab)], 6.3, align=PP_ALIGN.CENTER,
             wrap=False, name=f"bar value – {yr}")
        text(g, x - 6, base + 2, w + 12, 8, [(yr, 600, GREY)], 6, align=PP_ALIGN.CENTER, wrap=False,
             name=f"bar year – {yr}")
    hline(g, 90, base, 186, base, "9AA4B2", 0.8, name="bar baseline")

    # divider
    hline(g, 196, 22, 196, 99, LINE, 0.8, dash=MSO_LINE_DASH_STYLE.DASH, name="divider")

    # ---- right: govt. programme fit (NDMA / ISRO-NRSC)
    rows = [
        ("₹13,693 cr", RED, "Disaster Mitigation Fund\n(NDMF, 15th FC), 2021–26"),
        ("₹1,000 cr", ORANGE, "Landslide mitigation (NLRMP)\n15 states · approved Nov 2024"),
        ("₹150 cr", BLUE, "Glacial-lake flood mitigation\nSikkim, HP, UK, Arunachal"),
        ("NDEM", PURPLE, "ISRO-NRSC disaster geo-portal\nbuilt under NDMA guidance"),
    ]
    for i, (amt, col, desc) in enumerate(rows):
        y = 22 + i * 19.5
        text(g, 204, y, 58, 17, [(amt, 800, col)], 8.5, anchor=MSO_ANCHOR.MIDDLE, wrap=False, name=f"govt amount – {amt}")
        text(g, 262, y, 144, 17, [(desc, 500, NAVY)], 6.4, anchor=MSO_ANCHOR.MIDDLE, name=f"govt desc – {amt}")
    return g


# 1) full slide 4 with the placeholder replaced
prs = Presentation(D / "slide4_feasibility_viability.pptx")
s = prs.slides[0]
for sh in list(s.shapes):
    if sh.shape_id in (71, 141, 142):        # dashed box, 'MARKET & ADOPTION', placeholder text
        sh._element.getparent().remove(sh._element)
build(s.shapes)
prs.save(D / "slide4_feasibility_viability_market.pptx")

# 2) standalone component: same slide size and position, so paste lands in place
prs2 = Presentation(D / "slide4_feasibility_viability.pptx")
s2 = prs2.slides[0]
for sh in list(s2.shapes):
    sh._element.getparent().remove(sh._element)
build(s2.shapes)
prs2.save(D / "slide4_market_infographic.pptx")
print("ok")
